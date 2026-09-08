import os
import sys
import json
import re
import shutil
import difflib
from datetime import datetime
import openpyxl

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

from core.intelligent_matcher import phonetic_key

EXCEL_PATH = r"E:\Assembly.xlsx"

SHEET_DISTRICT_MAP = {
    'Rangareddy': 'Rangareddy',
    'Medchal malkajgiri': 'Medchal Malkajgiri',
    'Vikarabad': 'Vikarabad',
    'Nagarkarnul': 'Nagarkurnool',
    'MBNR': 'Mahabubnagar',
    'wanaparthy': 'Wanaparthy',
    'Jogulamba Gadwal': 'Jogulamba Gadwal',
    'Narayanpet': 'Narayanapet',
    'Nalgonda': 'Nalgonda',
    'Suryapet': 'Suryapet',
    'yadadri': 'Yadadri Bhuvanagiri',
    'Warangal': 'Warangal',
    'hanumakonda': 'Hanumakonda',
    'Mahabubabad': 'Mahabubabad',
    'mulugu': 'Mulugu',
    'Bhupalpally': 'Bhupapally',
    'kamareddy': 'Kamareddy',
    'Nzb': 'Nizamabad',
    'adb': 'Adilabad',
    'nirmal': 'Nirmal',
    'Asifabad': 'Asifabad',
    'Mancherial': 'Mancherial',
    'medak': 'Medak',
    'sangareddy': 'Sangareddy',
    'Siddipet': 'Siddipet',
    'Khammam': 'Khammam',
    'badradri': 'Kothagudem',
    'siricilla': 'Rajanna Siricilla',
    'Knr': 'Karimnagar',
    'peddapally': 'Peddapally',
    'jgtl': 'Jagitial',
    'janagaoan': 'Janagaon'
}

SPECIAL_VILLAGE_ALIASES = {
    ('Peddapally', 'janagaon'): 'jangoam',
    ('Rajanna Siricilla', 'pallemakta'): 'palle(makta)',
    ('Jogulamba Gadwal', 'm.r.chervu'): 'mogilravalcheruvu',
    ('Jogulamba Gadwal', 'mrchervu'): 'mogilravalcheruvu',
    ('Mulugu', 'padigapur(pattigorrevul'): 'padigapur',
    ('Adilabad', 'khairadatwa'): 'khairdatwa',
    ('Adilabad', 'taroda (srinath)'): 'tarada',
    ('Medchal Malkajgiri', 'anantharam'): 'anatharam',
    ('Siddipet', 'shivunipally'): 'sivampalle',
    ('Siddipet', 'arepally sj'): 'ankireddipalli',
    ('Sangareddy', 'kothur b'): 'kothur[b]',
}

def clean_v_name(text: str) -> str:
    if not text:
        return ""
    s = str(text).strip()
    m = re.findall(r'\(([A-Za-z0-9\s\.\-_]+)\)', s)
    if m and any(ord(c) > 127 for c in s):
        s = " ".join(m)

    s = s.lower()
    s = re.sub(r'\[[^\]]*\]|\([^)]*\)', ' ', s)
    for tok in [
        '413', '(v)', '(m)', ' vlg', ' vill', ' village', ' proper',
        ' khurd', ' kalan', ' buzurg', ' bk', ' kd', ' h/o', ' thanda', ' tanda'
    ]:
        s = s.replace(tok, ' ')
    s = re.sub(r'[\/\-\._,;:|*0-9]', ' ', s)
    tokens = s.split()
    noise = {'mandal', 'mdl', 'm', 'village', 'vill', 'vlg', 'v', 'dist', 'district'}
    tokens = [t for t in tokens if t not in noise]
    return ' '.join(tokens)

def sim_score(n1: str, n2: str) -> float:
    c1, c2 = clean_v_name(n1), clean_v_name(n2)
    if not c1 or not c2:
        return 0.0
    if c1 == c2:
        return 1.0
    p1, p2 = phonetic_key(c1), phonetic_key(c2)
    if p1 and p2 and p1 == p2:
        return 0.98
    w1, w2 = set(c1.split()), set(c2.split())
    if w1 and w2 and (w1.issubset(w2) or w2.issubset(w1)):
        return 0.92
    sm1 = difflib.SequenceMatcher(None, c1, c2).ratio()
    sm2 = difflib.SequenceMatcher(None, p1, p2).ratio() if p1 and p2 else 0.0
    return max(sm1, sm2)

def normalize_excel_extent(val):
    if val is None:
        return '0-00', 0.0
    s = str(val).strip()
    if not s or s.lower() in ['-', 'none', 'null', '0', '0.0', '0-00']:
        return '0-00', 0.0

    if '-' in s:
        parts = s.split('-')
        try:
            ac = int(float(parts[0]))
            gt = int(float(parts[1]))
            if gt >= 40:
                ac += gt // 40
                gt = gt % 40
            acres_float = round(ac + (gt / 40.0), 3)
            if gt == 0:
                return f'{ac}', acres_float
            return f'{ac}-{gt:02d}', acres_float
        except Exception:
            return s, 0.0

    try:
        f = float(s)
        ac = int(f)
        rem = round(f - ac, 6)
        if rem == 0:
            return f'{ac}', float(ac)

        if '.' in s:
            dec_str = s.split('.')[1]
            if len(dec_str) == 1:
                gt = int(dec_str) * 10
            else:
                gt = int(dec_str[:2])
        else:
            gt = round(rem * 40)

        if gt >= 40:
            ac += gt // 40
            gt = gt % 40

        acres_float = round(ac + (gt / 40.0), 3)
        if gt == 0:
            return f'{ac}', acres_float
        return f'{ac}-{gt:02d}', acres_float
    except Exception:
        return s, 0.0

def backup_files():
    backup_dir = os.path.join(REPO_ROOT, "data", "backups")
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    for f in ["master_villages.json", "villages.json"]:
        src = os.path.join(REPO_ROOT, "data", f)
        if os.path.exists(src):
            dst = os.path.join(backup_dir, f"{f}.{ts}.bak")
            shutil.copy2(src, dst)
            print(f"Backed up {f} -> {os.path.basename(dst)}")

def main():
    print("=" * 70)
    print("VILLAGE MASTER EXTENT UPDATE FROM EXCEL")
    print(f"Excel Source: {EXCEL_PATH}")
    print("=" * 70)

    if not os.path.exists(EXCEL_PATH):
        print(f"Error: Excel file not found at {EXCEL_PATH}")
        return

    backup_files()

    mv_path = os.path.join(REPO_ROOT, "data", "master_villages.json")
    v_path = os.path.join(REPO_ROOT, "data", "villages.json")

    with open(mv_path, "r", encoding="utf-8") as f:
        master_villages = json.load(f)

    with open(v_path, "r", encoding="utf-8") as f:
        operational_villages = json.load(f)

    print(f"Loaded {len(master_villages)} master villages and {len(operational_villages)} operational villages.")

    master_by_district = {}
    for mv in master_villages:
        d = mv.get("district_name", "").strip()
        if d not in master_by_district:
            master_by_district[d] = []
        master_by_district[d].append(mv)

    print("Loading Excel workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    updates_by_mv_id = {}
    district_stats = {}

    for sheet_name, std_district in SHEET_DISTRICT_MAP.items():
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet {sheet_name} not found in workbook.")
            continue

        ws = wb[sheet_name]
        db_district_vils = master_by_district.get(std_district, [])

        excel_rows = []
        curr_mandal = ""

        for r in range(1, ws.max_row + 1):
            m_val = ws.cell(row=r, column=4).value
            v_val = ws.cell(row=r, column=5).value
            e_val = ws.cell(row=r, column=6).value
            row_str = " ".join(str(ws.cell(row=r, column=c).value or "") for c in range(1, 10)).lower()

            if any(k in row_str for k in [
                "total", "sl.no", "resurvey report", "urbanized with vast", "nil"
            ]):
                continue

            if m_val and str(m_val).strip() and not str(m_val).strip().isdigit():
                curr_mandal = str(m_val).strip()

            if v_val and str(v_val).strip() and not str(v_val).strip().isdigit() and not str(v_val).strip().lower().startswith("total"):
                excel_rows.append((r, curr_mandal, str(v_val).strip(), e_val))

        matched_count = 0
        unmatched_rows = []

        for r_num, m_ex, v_ex, e_ex in excel_rows:
            v_ex_clean = clean_v_name(v_ex)
            v_ex_low = v_ex.lower()

            alias_target = SPECIAL_VILLAGE_ALIASES.get((std_district, v_ex_low))
            if alias_target:
                m_candidates = [mv for mv in db_district_vils if alias_target in mv.get("village_name", "").lower()]
            else:
                m_candidates = [mv for mv in db_district_vils if sim_score(m_ex, mv.get("mandal_name", "")) >= 0.5]

            best_mv = None
            best_score = 0.0

            for mv in m_candidates:
                score = sim_score(v_ex, mv.get("village_name", ""))
                if score > best_score:
                    best_score = score
                    best_mv = mv

            if best_score < 0.7:
                for mv in db_district_vils:
                    v_sc = sim_score(v_ex, mv.get("village_name", ""))
                    m_sc = sim_score(m_ex, mv.get("mandal_name", ""))
                    combined = v_sc * 0.8 + m_sc * 0.2
                    if v_sc >= 0.85:
                        combined = max(combined, v_sc)
                    if combined > best_score:
                        best_score = combined
                        best_mv = mv

            if best_mv and best_score >= 0.60:
                ext_str, ext_flt = normalize_excel_extent(e_ex)
                vid = best_mv["id"]

                if vid not in updates_by_mv_id or (ext_flt > 0 and updates_by_mv_id[vid]["extent_acres_float"] == 0):
                    updates_by_mv_id[vid] = {
                        "extent_existing_record": ext_str,
                        "extent_raw": ext_str,
                        "extent_acres_float": ext_flt,
                        "is_picked_for_resurvey": True,
                        "matched_excel_village": v_ex,
                        "matched_excel_mandal": m_ex,
                        "db_village_name": best_mv.get("village_name"),
                        "db_mandal_name": best_mv.get("mandal_name"),
                        "district": std_district,
                        "score": best_score
                    }
                matched_count += 1
            else:
                unmatched_rows.append((r_num, m_ex, v_ex, e_ex, best_score, best_mv.get("village_name") if best_mv else "None"))

        district_stats[std_district] = {
            "excel_rows": len(excel_rows),
            "matched": matched_count,
            "unmatched": len(unmatched_rows)
        }

    total_master_updated = 0
    for idx, mv in enumerate(master_villages):
        vid = mv.get("id")
        if vid in updates_by_mv_id:
            upd = updates_by_mv_id[vid]
            master_villages[idx]["extent_existing_record"] = upd["extent_existing_record"]
            master_villages[idx]["extent_raw"] = upd["extent_raw"]
            master_villages[idx]["extent_acres_float"] = upd["extent_acres_float"]
            master_villages[idx]["is_picked_for_resurvey"] = True
            if master_villages[idx].get("resurvey_phase") in [None, "", "Not Picked (Subsequent Phase)"]:
                master_villages[idx]["resurvey_phase"] = "Phase 1 (Active Resurvey)"
            total_master_updated += 1

    total_operational_updated = 0
    mv_id_to_extent = {vid: upd for vid, upd in updates_by_mv_id.items()}

    for idx, ov in enumerate(operational_villages):
        ov_id = ov.get("id")
        d_name = ov.get("district_name", "")
        v_name = ov.get("village_name", "")
        m_name = ov.get("mandal_name", "")

        matched_upd = None
        if ov_id in mv_id_to_extent:
            matched_upd = mv_id_to_extent[ov_id]
        else:
            best_sc = 0.0
            for vid, upd in updates_by_mv_id.items():
                if upd["district"] == d_name:
                    v_sc = sim_score(v_name, upd["db_village_name"])
                    m_sc = sim_score(m_name, upd["db_mandal_name"])
                    comb = v_sc * 0.7 + m_sc * 0.3
                    if v_sc >= 0.85:
                        comb = max(comb, v_sc)
                    if comb > best_sc and comb >= 0.60:
                        best_sc = comb
                        matched_upd = upd

        if matched_upd:
            operational_villages[idx]["extent_existing_record"] = matched_upd["extent_existing_record"]
            operational_villages[idx]["extent_raw"] = matched_upd["extent_raw"]
            operational_villages[idx]["extent_acres_float"] = matched_upd["extent_acres_float"]
            total_operational_updated += 1

    paths_to_save = [
        (mv_path, master_villages),
        (os.path.join(REPO_ROOT, "api", "data", "master_villages.json"), master_villages),
        (v_path, operational_villages),
        (os.path.join(REPO_ROOT, "api", "data", "villages.json"), operational_villages)
    ]

    for fpath, data in paths_to_save:
        os.makedirs(os.path.dirname(fpath), exist_ok=True)
        with open(fpath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"Saved: {fpath} ({len(data)} records)")

    print("\n" + "=" * 70)
    print("UPDATE SUMMARY BY DISTRICT")
    print(f"{'District':<24} | {'Excel Rows':<11} | {'Matched':<9} | {'Unmatched':<9}")
    print("-" * 70)
    total_ex = 0
    total_m = 0
    total_u = 0
    for d, s in sorted(district_stats.items()):
        print(f"{d:<24} | {s['excel_rows']:<11} | {s['matched']:<9} | {s['unmatched']:<9}")
        total_ex += s['excel_rows']
        total_m += s['matched']
        total_u += s['unmatched']
    print("-" * 70)
    print(f"{'TOTAL':<24} | {total_ex:<11} | {total_m:<9} | {total_u:<9}")
    print("=" * 70)
    print(f"Total Unique Master Villages Updated: {total_master_updated}")
    print(f"Total Operational Villages Updated: {total_operational_updated}")
    print("=" * 70)

if __name__ == "__main__":
    main()
